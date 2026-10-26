import os
import re
import io
import unicodedata
from typing import List, Dict, Any 
from pypdf import PdfReader
from docx import Document
import tiktoken


class DocumentProcessor:
    def __init__(self, encoding_name: str = "cl100k_base"):
        self.tokenizer = tiktoken.get_encoding(encoding_name)

    def count_tokens(self, text: str) -> int:
        return len(self.tokenizer.encode(text))

    def clean_page_text(self, text: str) -> str:
        """Làm sạch kỹ văn bản từng trang: loại bỏ khoảng trắng thừa, số trang chân trang, rác"""
        text = unicodedata.normalize('NFC', text)
        
        # Loại bỏ số trang ở đầu/chân trang
        text = re.sub(r'\b(Trang|Page)\s*\d+(\s*(?:của|of)\s*\d+)?\b', '', text, flags=re.IGNORECASE)
        text = re.sub(r'^\s*-\s*\d+\s*-\s*$', '', text, flags=re.MULTILINE)
        
        text = re.sub(r'(?<=[à-ỹa-z])\s+(?=[à-ỹa-z]\b)', '', text, flags=re.IGNORECASE)
        text = re.sub(r'\b([à-ỹa-z])\s+([à-ỹa-z])\b', r'\1\2', text, flags=re.IGNORECASE)
        
        # Chuẩn hóa các khoảng trắng thừa thành 1 khoảng trắng đơn
        text = re.sub(r'[ \t]+', ' ', text)
        return text.strip()

    # --- CÁC HÀM XỬ LÝ THEO FILE PATH (Giữ nguyên để tương thích ngược) ---
    def extract_text_from_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        try:
            reader = PdfReader(file_path)
            return self._extract_pdf_from_reader(reader)
        except Exception as e:
            raise ValueError(f"Lỗi khi đọc file PDF: {str(e)}")

    def extract_text_from_docx(self, file_path: str) -> List[Dict[str, Any]]:
        try:
            doc = Document(file_path)
            return self._extract_docx_from_doc(doc)
        except Exception as e:
            raise ValueError(f"Lỗi khi đọc file DOCX: {str(e)}")

    # --- CÁC HÀM MỚI: XỬ LÝ TRỰC TIẾP TỪ STREAM / BYTES (Dành cho Celery qua RAM) ---
    def extract_text_from_pdf_stream(self, file_stream: io.BytesIO) -> List[Dict[str, Any]]:
        try:
            reader = PdfReader(file_stream)
            return self._extract_pdf_from_reader(reader)
        except Exception as e:
            raise ValueError(f"Lỗi khi đọc stream PDF: {str(e)}")

    def extract_text_from_docx_stream(self, file_stream: io.BytesIO) -> List[Dict[str, Any]]:
        try:
            doc = Document(file_stream)
            return self._extract_docx_from_doc(doc)
        except Exception as e:
            raise ValueError(f"Lỗi khi đọc stream DOCX: {str(e)}")

    # --- HÀM HỖ TRỢ NỘI BỘ TRÍCH XUẤT ---
    def _extract_pdf_from_reader(self, reader: PdfReader) -> List[Dict[str, Any]]:
        pages_content = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            cleaned_text = self.clean_page_text(text)
            if cleaned_text:
                pages_content.append({
                    "page_number": index + 1,
                    "content": cleaned_text
                })
        return pages_content

    def _extract_docx_from_doc(self, doc: Document) -> List[Dict[str, Any]]:
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    full_text.append(row_text)
        
        combined_text = " ".join(full_text)
        cleaned_text = self.clean_page_text(combined_text)
        
        return [{
            "page_number": 1,
            "content": cleaned_text
        }]

    def chunk_text(self, pages: List[Dict[str, Any]], min_chars: int = 500, max_chars: int = 1000) -> List[Dict[str, Any]]:
        """Cắt chunk và duy trì tiêu đề chương/mục xuyên suốt"""
        chunks = []
        chunk_index = 0
        current_section_title = "Phần mở đầu / Tài liệu chung"

        heading_regex = re.compile(r'^(CHƯƠNG\s+\d+|PHẦN\s+\d+|\d+[\.\)]\s+[A-ZÀ-Ỹ].*|\d+\.\d+\s+[A-ZÀ-Ỹ].*)', re.IGNORECASE)
        separators = ["\n\n", "\n", ". ", ", ", " "]

        all_lines = []
        for page in pages:
            page_num = page["page_number"]
            text = page["content"]
            lines = text.split('\n')
            for line in lines:
                line_str = line.strip()
                if line_str:
                    all_lines.append((page_num, line_str))

        structured_sections = []
        current_block = []
        current_page_num = 1

        for page_num, line_str in all_lines:
            current_page_num = page_num
            if heading_regex.match(line_str):
                if current_block:
                    structured_sections.append((current_section_title, "\n".join(current_block), current_page_num))
                    current_block = []
                current_section_title = line_str 
                current_block.append(line_str)
            else:
                current_block.append(line_str)

        if current_block:
            structured_sections.append((current_section_title, "\n".join(current_block), current_page_num))

        for sec_title, sec_content, page_num in structured_sections:
            sub_chunks = self._recursive_split(sec_content, separators, max_chars, min_chars)
            for sub_text in sub_chunks:
                final_content = sub_text.strip()
                if len(final_content) > 20:
                    chunks.append({
                        "chunk_index": chunk_index,
                        "content": final_content,
                        "page_number": page_num,
                        "token_count": self.count_tokens(final_content),
                        "section_title": sec_title 
                    })
                    chunk_index += 1

        return chunks

    def _recursive_split(self, text: str, separators: List[str], max_chars: int, min_chars: int) -> List[str]:
        if len(text) <= max_chars:
            return [text]

        current_separator = separators[0]
        next_separators = separators[1:] if len(separators) > 1 else [" "]

        splits = text.split(current_separator)
        good_splits = []
        current_chunk = []
        current_length = 0

        for part in splits:
            part_len = len(part) + len(current_separator)
            if current_length + part_len <= max_chars:
                current_chunk.append(part)
                current_length += part_len
            else:
                if current_chunk:
                    joined = current_separator.join(current_chunk)
                    if len(joined) >= min_chars or not next_separators:
                        good_splits.append(joined)
                    else:
                        sub = self._recursive_split(joined, next_separators, max_chars, min_chars)
                        good_splits.extend(sub)
                    current_chunk = [part]
                    current_length = len(part)
                else:
                    sub = self._recursive_split(part, next_separators, max_chars, min_chars)
                    good_splits.extend(sub)

        if current_chunk:
            joined = current_separator.join(current_chunk)
            if len(joined) >= min_chars or not next_separators:
                good_splits.append(joined)
            else:
                sub = self._recursive_split(joined, next_separators, max_chars, min_chars)
                good_splits.extend(sub)

        return good_splits

    def process_file(self, file_path: str, file_extension: str) -> List[Dict[str, Any]]:
        ext = file_extension.lower()
        if ext == ".pdf":
            pages = self.extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            pages = self.extract_text_from_docx(file_path)
        elif ext == ".txt":
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
            except UnicodeDecodeError:
                with open(file_path, "r", encoding="latin-1") as f:
                    content = f.read()
            cleaned_content = self.clean_page_text(content)
            pages = [{"page_number": 1, "content": cleaned_content}]
        else:
            raise ValueError(f"Định dạng tệp không được hỗ trợ: {ext}")

        return self.chunk_text(pages)

    # --- HÀM XỬ LÝ STREAM MỚI CHO CELERY ---
    def process_stream(self, file_stream: io.BytesIO, file_extension: str) -> List[Dict[str, Any]]:
        ext = file_extension.lower()
        if ext == ".pdf":
            pages = self.extract_text_from_pdf_stream(file_stream)
        elif ext in [".docx", ".doc"]:
            pages = self.extract_text_from_docx_stream(file_stream)
        elif ext == ".txt":
            try:
                content = file_stream.read().decode("utf-8")
            except UnicodeDecodeError:
                file_stream.seek(0)
                content = file_stream.read().decode("latin-1")
            cleaned_content = self.clean_page_text(content)
            pages = [{"page_number": 1, "content": cleaned_content}]
        else:
            raise ValueError(f"Định dạng tệp không được hỗ trợ: {ext}")

        return self.chunk_text(pages)